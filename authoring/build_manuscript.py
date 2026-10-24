"""Build the StreetLux paper from editable prose, saved results and equations."""
from pathlib import Path
import argparse
import copy
import json
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from make_decomposition import calculate

def load(path):
    return json.loads(path.read_text(encoding='utf-8'))

def table_data(key):
    if key == 'street_route':
        route = load(RESULTS / 'route.json')
        rows = [[s['name'].replace('Kazi Nazrul arterial', 'Wide arterial').replace('Motijheel office canyon', 'Office canyon'), f"{s['length_m']:.0f}", f"{s['width_m']:g}", f"{s['left_height_m']:g} / {s['right_height_m']:g}", f"{s['azimuth_deg']:g}", f"{s['eye_offset_m']:g}"] for s in route['sections']]
        return 'Table 1. Declared route geometry. All distances are in metres; azimuth is clockwise from north. Positive eye offsets are to the right of the centreline in the direction of travel.', ['Section', 'Length', 'Width', 'Left / right height', 'Azimuth °', 'Eye offset'], rows, [1.7, .65, .6, 1.05, .8, .7]
    if key == 'street_modes':
        data = load(RESULTS / 'modes.json')['rows']
        names = {'walk': 'Walking', 'rickshaw': 'Rickshaw', 'cng': 'CNG auto-rickshaw', 'car_clear': 'Clear-glazed car', 'car_tinted': 'Bronze-glazed car', 'bus': 'Bus window seat', 'car_film': 'Car with film'}
        rows = [[names[d['key']], f"{d['summary']['duration_min']:.2f}", f"{d['summary']['mean_medi_travel_lx']:.1f}", f"{d['summary']['dose_travel_lxh']:.1f}", f"{d['summary']['mean_der']:.3f}", f"{d['summary']['heading_spread']['ratio_max_min']:.2f}"] for d in data]
        return 'Table 2. Mode comparison. DER is the time-mean melanopic daylight efficacy ratio; heading ratio uses fixed world headings.', ['Mode', 'Time min', 'mEDI lx', 'Dose lx h', 'DER', 'Heading ratio'], rows, [1.6, .7, .8, .85, .6, .95]
    if key == 'street_shares':
        data = calculate(load(RESULTS / 'modes.json'))
        columns = ['duration_first_pct', 'rate_first_pct', 'symmetric_duration_pct', 'symmetric_rate_pct']
        rows = [[r['label']] + [f'{r[c]:.1f}' for c in columns] for r in data]
        return 'Table 3. Dose-deficit attribution relative to walking. Sequential columns report the duration share under each order. The symmetric duration and rate shares sum to 100%. All values are percentages.', ['Mode', 'Duration first', 'Rate first', 'Symmetric duration', 'Symmetric rate'], rows, [1.8, 1.175, 1.175, 1.175, 1.175]
    if key == 'street_designs':
        data = load(RESULTS / 'designs.json')['rows']
        names = {'pale_facades': 'Whitewashed façades', 'setback_low': 'Height cap 18 m', 'widen_18': 'Minimum width 18 m', 'pale_paving': 'Asphalt to pale concrete', 'baseline': 'Baseline', 'canopy_40': 'Foliage cover 40%', 'canopy_70': 'Foliage cover 70%'}
        rows = [[names[d['key']], f"{d['summary']['mean_medi_travel_lx']:.1f}", f"{d['summary']['dose_travel_lxh']:.1f}", f"{d['delta_vs_baseline_pct']:+.2f}"] for d in data]
        return 'Table 4. Independent walking interventions. Changes are relative to the baseline and are identical for mean mEDI and dose because duration is fixed.', ['Scenario', 'mEDI lx', 'Dose lx h', 'Change %'], rows, [2.65, 1, 1.15, .9]
    raise ValueError(f'Unknown table: {key}')

def shade_cell(cell,fill):
    pr=cell._tc.get_or_add_tcPr();el=OxmlElement('w:shd');el.set(qn('w:fill'),fill);pr.append(el)

def add_table(doc,key):
    caption,headers,rows,widths=table_data(key)
    if caption:
        p=doc.add_paragraph(caption,'Caption');p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    total=sum(widths);widths=[6.5*w/total for w in widths]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
    pr.append(b)
    for i,txt in enumerate(headers):t.rows[0].cells[i].text=txt
    h=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(h)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for ri,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            cell._tc.get_or_add_tcPr().append(mar)
            shade_cell(cell,'DDE8EE' if ri==0 else ('F6F8F9' if ri%2==0 else 'FFFFFF'))
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(headers)==2 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9.5);r.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--results', type=Path, default=ROOT / 'results')
parser.add_argument('--figures', type=Path, default=ROOT / 'figures')
parser.add_argument('--output', type=Path, default=ROOT / 'manuscript/StreetLux_Unified_Manuscript.docx')
args = parser.parse_args()
RESULTS, FIGURES = args.results, args.figures
text = (ROOT / 'manuscript.md').read_text(encoding='utf-8')
abstract = re.search(r'## Abstract\n\n(.*?)\n\nKeywords:', text, re.S).group(1)
maths = [p for p in Document(ROOT / 'authoring/equations.docx').paragraphs if p._p.xpath('.//m:oMath')]
assert len(maths) == 3

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.8)
sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.35)
styles=doc.styles
for border in styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
for name,size in [('Normal',11),('Title',18),('Heading 1',13),('Heading 2',11.5),('Caption',9.5)]:
    st=styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(6)
    rf=st.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in list(rf.attrib):
        if 'Theme' in attr:del rf.attrib[attr]
    for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
styles['Normal'].paragraph_format.line_spacing=1.12
styles['Normal'].paragraph_format.widow_control=True
styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(12)
for name in ['Heading 1','Heading 2']:
    styles[name].font.bold=True
    styles[name].paragraph_format.keep_with_next=True
    styles[name].paragraph_format.space_before=Pt(10)
styles['Caption'].font.italic=False
styles['Caption'].font.bold=False
styles['Caption'].paragraph_format.line_spacing=1.05
header=sec.header.paragraphs[0]
header.text='StreetLux  |  Research manuscript'
for r in header.runs:r.font.name='Times New Roman';r.font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

references=False
for block in re.split(r'\n\s*\n',text):
    block=block.strip()
    if not block:continue
    if block.startswith('{{fig:'):
        src,cap=block[6:-2].split('|',1)
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
        p.add_run().add_picture(str(FIGURES/(src+'.png')),width=Inches(6.5))
        for el in p._p.xpath('.//wp:docPr'):el.set('descr',cap)
        cp=doc.add_paragraph(cap,'Caption');cp.paragraph_format.space_after=Pt(9)
    elif block.startswith('{{table:'):
        add_table(doc,block[8:-2])
        # All four tables are small enough to stay intact on one page.
        for row in doc.tables[-1].rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
    elif block.startswith('{{eq:'):
        i=['metrology','sequential','symmetric'].index(block[5:-2])
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for math in maths[i]._p.xpath('.//m:oMath'):
            p._p.append(copy.deepcopy(math))
        p.add_run('     ('+str(i+1)+')')
        p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(8)
    elif block.startswith('# '):doc.add_paragraph(block[2:],'Title')
    elif block.startswith('### '):doc.add_paragraph(block[4:],'Heading 2')
    elif block.startswith('## '):
        title=block[3:];doc.add_paragraph(title,'Heading 1')
        if title=='References':references=True
    else:
        p=doc.add_paragraph(block)
        if references:
            p.paragraph_format.left_indent=Inches(.18)
            p.paragraph_format.first_line_indent=Inches(-.18)
            p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_together=True
            p.paragraph_format.keep_with_next=not block.startswith('Shapley LS.')
            for r in p.runs:r.font.size=Pt(10)
        elif block.startswith('Keywords:'):
            for r in p.runs:r.font.size=Pt(10)

doc.core_properties.title=text.splitlines()[0][2:]
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.subject='Computational study of melanopic daylight exposure during travel'
doc.core_properties.comments=''
target = args.output
target.parent.mkdir(parents=True, exist_ok=True)
doc.save(target)
print(json.dumps({'docx':str(target),'abstract_words':len(abstract.split()),'text_words':len(text.split()),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'native_equations':len(maths)},indent=2))
assert len(abstract.split())<=250
assert len(doc.tables)==4 and len(doc.inline_shapes)==4
